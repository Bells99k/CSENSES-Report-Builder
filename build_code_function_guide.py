from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, "/private/tmp/codex-docx-deps")

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "CSENSES_Code_Blocks_and_Functions_Guide.docx"
BLACK = "000000"


def font(run, name="Arial", size=10.5, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(BLACK)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def remove_style_border(style):
    ppr = style._element.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    font(run, size=8.5)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def set_left_border(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "8")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), BLACK)
    pbdr.append(left)
    ppr.append(pbdr)


def add_code(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    set_left_border(p)
    lines = code.strip("\n").split("\n")
    for index, line in enumerate(lines):
        r = p.add_run(line)
        font(r, name="Consolas", size=8.3)
        if index < len(lines) - 1:
            r.add_break()


def add_intro(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    font(r, size=10.5)


def add_function(doc, names, description):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    name_list = names if isinstance(names, (list, tuple)) else [names]
    for index, name in enumerate(name_list):
        if index:
            sep = p.add_run(", ")
            font(sep, size=9.5)
        r = p.add_run(name)
        font(r, name="Consolas", size=9.2, bold=True)
    r = p.add_run(" — " + description)
    font(r, size=9.7)


def add_block(doc, title, purpose, code=None, functions=None):
    doc.add_heading(title, level=2)
    add_intro(doc, purpose)
    if code:
        add_code(doc, code)
    for names, description in functions or []:
        add_function(doc, names, description)


def configure_document(doc):
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    section.different_first_page_header_footer = False

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15

    style_tokens = {
        "Title": (20, True, 0, 4),
        "Subtitle": (11, False, 0, 12),
        "Heading 1": (15, True, 14, 6),
        "Heading 2": (12, True, 10, 4),
        "Heading 3": (10.5, True, 7, 3),
    }
    for style_name, (size, bold, before, after) in style_tokens.items():
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.0
        remove_style_border(style)

    footer = section.footer
    add_page_field(footer.paragraphs[0])


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("CSENSES Report Builder: Code Blocks and Functions")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Grouped explanation of the current codebase · commit e15cf32 · September 23, 2026")

    add_intro(
        doc,
        "How to read this guide: each Block Group covers one responsibility. Within it, each Block explains a related section of code, shows the core pattern, and identifies the functions that do the work. Small helpers are grouped when they serve the same purpose.",
    )

    doc.add_heading("Block Group 1 — Application state and fixed definitions (app.js: 1–680)", level=1)
    add_block(
        doc,
        "Block 1.1 — Shared state and DOM references",
        "This is the application's memory and its connection to the HTML. Nearly every later function reads from or writes to these two objects.",
        """// abridged
const state = { template, rows, sensorCatalog, clusterCatalog,
  comparisonLocations, customClusters, noteHtml, apiAbortController, ... };

const els = {
  location: document.getElementById("locationName"),
  calendarMetric: document.getElementById("calendarMetric"), ...
};""",
        [
            ("state", "Holds current template, loaded rows, catalogs, map state, notes, images, comparisons, and active API-load information."),
            ("els", "Caches important HTML elements once so renderers and event handlers can reuse them."),
        ],
    )
    add_block(
        doc,
        "Block 1.2 — Constants, catalogs, standards, and sample data",
        "These definitions provide API endpoints, timeouts, map settings, sensor metadata, threshold bands, image paths, and fallback data. They are configuration rather than executable functions.",
        """const apiRequestTimeoutMs = 100000;
const maxConcurrentSensorRequests = 4;
const metricStandards = { composite: {...}, air: {...}, pm10: {...},
  heat: {...}, noise: {...} };
const sampleRows = ...;""",
        [
            (["builtInSensorCatalog", "builtInClusterCatalog"], "Fallback location metadata used before, or instead of, remote catalogs."),
            (["metricStandards", "heatIndexBands"], "Defines labels, numeric ranges, colors, health effects, and recommendations used by the report."),
            (["pictureAssets", "sampleRows"], "Provides bundled photos and immediately available demonstration readings."),
        ],
    )

    doc.add_heading("Block Group 2 — Parsing and catalog loading (app.js: 681–956)", level=1)
    add_block(
        doc,
        "Block 2.1 — File parsers",
        "Converts CSV text and shapefile DBF bytes into JavaScript rows that can be normalized into the app's data structures.",
        """const rows = normalizeRows(parseCsv(csvText));
const metadata = parseDbf(arrayBuffer);""",
        [
            ("parseCsv(text)", "Parses quoted or unquoted CSV fields and returns a two-dimensional array."),
            ("parseDbf(arrayBuffer)", "Reads DBF headers, field definitions, and records used by the bundled shapefiles."),
            ("cleanSensorId(value)", "Normalizes sensor IDs so file, API, and UI identifiers compare reliably."),
        ],
    )
    add_block(
        doc,
        "Block 2.2 — Endpoint and identifier helpers",
        "Builds catalog/readings endpoint paths and reconciles display IDs, API location IDs, and fallback sensor names.",
        None,
        [
            (["sensorListUrl(namespace)", "sensorReadingsUrl(namespace)", "clusterReadingsUrl(namespace)"], "Returns the AQ or NU catalog/readings endpoint for the requested namespace."),
            (["clusterListUrl()", "clusterListUrls()"], "Returns the preferred cluster catalog endpoint and its local fallback path."),
            (["sensorDisplayId(kind, filterId)", "catalogKeyedByLocation(catalog)"], "Creates user-facing IDs and deduplicates catalog entries by location."),
            (["normalizeSensorName(value)", "fallbackSensorDetails(kind, filterId, name)"], "Normalizes names for matching and supplies built-in metadata when the API is incomplete."),
            ("legacyApiLocationIdForSensor(sensorOrSelection)", "Maps older sensor IDs to current API location IDs."),
        ],
    )
    add_block(
        doc,
        "Block 2.3 — Remote and local catalog loaders",
        "Prefers live API catalogs, but keeps built-in and DBF-derived metadata available as fallbacks.",
        """const remoteCatalog = await loadRemoteSensorCatalog();
state.sensorCatalog = remoteCatalog.length
  ? catalogKeyedByLocation(remoteCatalog)
  : catalogKeyedByLocation(builtInSensorCatalog);""",
        [
            ("normalizeRemoteSensorCatalog(payload, source)", "Converts AQ/NU catalog responses into the common sensor object shape."),
            ("loadRemoteSensorCatalog()", "Fetches AQ and NU sensor lists and combines usable results."),
            ("normalizeRemoteClusterCatalog(payload)", "Converts the cluster API response into ID/name records."),
            ("loadClusterCatalog()", "Loads the remote or local cluster list and stores it in state."),
            ("loadSensorCatalog()", "Uses the remote catalog when possible; otherwise parses the bundled DBF files."),
        ],
    )

    doc.add_heading("Block Group 3 — Normalized rows, dates, metrics, and locations (app.js: 957–1426)", level=1)
    add_block(
        doc,
        "Block 3.1 — Common row format",
        "Every source is converted to the same row shape. Downstream charts and summaries therefore do not need to know whether data came from sample rows, CSV, or the API.",
        """{
  date: "YYYY-MM-DD", cluster: "...", sensorId: "...",
  air: number|null, pm10: number|null,
  heat: number|null, noise: number|null
}""",
        [
            ("normalizeRows(csvRows)", "Finds accepted CSV column aliases and emits the common row shape."),
            (["cleanText(value)", "cleanDate(value)", "toNumber(value)", "roundNumber(value)"], "Sanitizes strings, ISO dates, nullable numeric values, and one-decimal display values."),
        ],
    )
    add_block(
        doc,
        "Block 3.2 — Period, metric, and status helpers",
        "Centralizes the selected month/day, units, labels, thresholds, and loading feedback.",
        None,
        [
            (["setDataStatus(message, tone)", "setDataLoadProgress(isLoading)"], "Updates the accessible status message and loading indicator."),
            (["monthInfo()", "monthDateRange()", "selectedDataDateRange()"], "Calculates calendar metadata and the date range required by the active template/aggregation."),
            ("thresholds()", "Reads the hidden report threshold inputs."),
            (["metricUnit(metric)", "metricDisplay(metric)", "metricDisplayHtml(metric)", "metricTextHtml(value)"], "Provides plain-text or HTML-safe metric labels and units."),
            ("allSensorLocationsLabel(metric)", "Builds the Template 3 label used when all matching sensors are included."),
        ],
    )
    add_block(
        doc,
        "Block 3.3 — Filtering and daily averaging",
        "Selects the active period and location, then averages multiple sensor rows into one row per date before statistics or rendering.",
        """state.rows
  .filter(row => rowMatchesMonth(row, start, end))
  .filter(row => rowMatchesLocation(row, selection))
  -> averageRowsByDate(rows);""",
        [
            (["filteredRows()", "selectedMetricRows()"], "Returns daily averaged rows for the active month/location, optionally requiring the selected metric."),
            (["rowHasMetricValue(row, metric)", "rowMatchesMonth(row, startDate, endDate)"], "Tests metric presence and date inclusion."),
            (["monthRowsForCluster(cluster)", "monthRowsForLocation(value)"], "Returns daily rows for a specific comparison target."),
            ("averageRowsByDate(rows)", "Groups rows by date and averages each available metric."),
            ("averageMetric(metric)", "Calculates the selected period's mean for one metric."),
        ],
    )
    add_block(
        doc,
        "Block 3.4 — Location options and selection resolution",
        "Builds the available sensor/cluster/custom-cluster options and converts a select value into a complete selection object.",
        None,
        [
            (["clusterOptions()", "clusterOptionsFromRows()", "clusterOptionsFromCatalog()"], "Builds cluster options from loaded data and the catalog."),
            (["sensorOptionsFromRows()", "sensorOptionsFromCatalog()", "sensorLocationOptions()"], "Builds sensor options from data and metadata, filtered by metric compatibility."),
            (["baseLocationOptions()", "customClusterOptions(metric)", "locationOptions()", "comparisonLocationOptions()"], "Combines normal and custom locations for primary and comparison selectors."),
            (["selectedHazardMetric()", "metricHasValue(row, metric)", "catalogSensorMatchesMetric(sensor, metric)"], "Identifies the active metric and compatible data/sensors."),
            (["locationValue(kind, id)", "parseLocationValue(value)", "selectedLocation(value, options)"], "Encodes, decodes, and resolves selector values."),
            (["locationDisplay(value)", "reportLocationDisplay(value)"], "Returns short UI text or report-ready location text."),
            ("rowMatchesLocation(row, selection)", "Checks whether a row belongs to a sensor, predefined cluster, or custom cluster."),
            (["resolvePreferredLocation(...) ", "equivalentSensorLocationValue(...)"], "Preserves a valid selection when catalogs or metric filters change."),
            (["defaultSensorFilterIdForMetric(metric)", "defaultPreferredLocationValue()", "metricChangePreferredLocationValue()", "defaultLocationValue(options)"], "Chooses safe default selections."),
            (["updateClusterOptions(preferredCluster)", "syncVisibleLocationSelects(value)", "populateClusterSelect(select, options)"], "Rebuilds the hidden canonical selector and its visible sensor/cluster controls."),
        ],
    )

    doc.add_heading("Block Group 4 — Custom clusters, search, and comparisons (app.js: 1427–1933)", level=1)
    add_block(
        doc,
        "Block 4.1 — Custom-cluster builder",
        "Maintains draft sensor lists, validates a cluster name, loads member data, and creates a reusable in-memory custom cluster.",
        """draft sensor IDs -> load selected data -> createCustomCluster()
  -> state.customClusters -> locationOptions() -> report""",
        [
            (["customClusterSensorOptions()", "clusterBuilderConfig(context)"], "Returns compatible sensors and the correct set of controls for the primary or comparison builder."),
            (["setCustomClusterStatus(...) ", "setCustomClusterNameError(...)"], "Shows validation/loading state."),
            (["renderGroupedOptions(...) ", "renderCustomClusterBuilder(context)", "renderCustomClusterBuilders()"], "Renders grouped sensor choices, selected-member chips, and button states."),
            ("addCustomClusterDraftSensor(context)", "Adds the chosen sensor once to the active draft."),
            ("createCustomCluster(context)", "Validates the draft, stores a new custom cluster, and selects it."),
            ("startNewCustomCluster(context)", "Clears builder state for another cluster."),
        ],
    )
    add_block(
        doc,
        "Block 4.2 — Sensor, cluster, and comparison search",
        "Filters existing options by user text and connects search results back to the underlying select controls.",
        None,
        [
            (["matchingLocationOptions(query)", "renderSensorSearchResults(options)", "searchLocationOptions()"], "Finds, displays, and selects sensor search matches."),
            (["hideSensorSearchResults()", "selectLocationSearchResult(value, label)"], "Closes results and commits a sensor match."),
            (["renderClusterSearchResults(options)", "selectClusterSearchResult(value, label)", "hideClusterSearchResults()"], "Provides the same workflow for predefined clusters."),
            (["matchingComparisonOptions(query)", "renderComparisonSearchResults(options)", "addFirstComparisonSearchResult()"], "Finds comparison candidates and supports Enter-key selection."),
            (["hideComparisonSearchResults()", "addComparisonLocation(value)"], "Closes the result list and adds a unique comparison."),
        ],
    )
    add_block(
        doc,
        "Block 4.3 — Comparison defaults and series",
        "Keeps comparison selections valid and converts each location into a dated series for Template 2 and Template 3.",
        None,
        [
            (["primaryComparisonLocationValue(...) ", "defaultComparisonLocations(...)"], "Selects sensible comparison defaults around the primary location."),
            (["updateComparisonLocationOptions(...) ", "renderComparisonLocationOptions()", "renderComparisonSelectOptions()"], "Refreshes compatible comparison choices."),
            (["comparisonOption(value)", "renderComparisonSelected()", "selectedComparisonLocations()"], "Resolves and displays currently selected comparisons."),
            ("comparisonSeries()", "Produces named/color-coded daily series used by the trend chart and comparison summaries."),
        ],
    )

    doc.add_heading("Block Group 5 — Statistics, standards, calendar, and trend chart (app.js: 1934–2454)", level=1)
    add_block(
        doc,
        "Block 5.1 — Summary calculations",
        "Computes threshold counts, peaks, band frequencies, and comparison outcomes from daily averaged rows.",
        """daily rows -> metricSummaryStats()
  -> exceedanceDays, maxValue, peakRow, topBand""",
        [
            ("exceedanceStats()", "Calculates per-metric threshold counts and maximums for the selected rows."),
            ("heatSummaryStats(rows)", "Counts heat caution-or-higher days, maximum heat, and danger/extreme-danger days."),
            ("metricSummaryStats(rows, metric)", "Produces the generic summary cards for any exposed metric."),
        ],
    )
    add_block(
        doc,
        "Block 5.2 — Standard-band classification",
        "Maps a reading to the correct category, label, color, and accessible text.",
        None,
        [
            (["rowSeverity(row, metric, limits)", "heatColor(severity)", "heatLabel(severity)"], "Calculates normalized composite severity and legacy gradient labels/colors."),
            (["heatIndexBand(value)", "standardBand(metric, row, severity)"], "Finds the matching Heat Index or general standards band."),
            (["calendarColor(...) ", "calendarLabel(...) ", "calendarBandClass(...)"], "Converts a day's band into calendar presentation values."),
            (["renderStandardsGraphic(graphic, standard)", "renderStandards(metric)", "darkTextOn(color)"], "Builds the standards legend and selects readable foreground text."),
        ],
    )
    add_block(
        doc,
        "Block 5.3 — Calendar renderer",
        "Builds one grid cell per calendar day, looks up that date's averaged row, and applies the selected standards band.",
        """for (day = 1; day <= month.days; day++) {
  const row = stats.byDate.get(date);
  cell.style.background = calendarColor(row, metric, severity);
}""",
        [
            ("renderCalendar()", "Creates weekday headers, blank leading cells, dated value cells, labels, and no-data cells."),
            ("formatCellValue(row, metric, severity)", "Formats the daily value shown inside a calendar cell."),
        ],
    )
    add_block(
        doc,
        "Block 5.4 — Trend chart and tooltip",
        "Draws standards bands and multiple daily series on a canvas, then resolves pointer position back to a day/value tooltip.",
        None,
        [
            (["trendChartScale(values, threshold, metric)", "trendValueToY(value, scale, padding, plotH)"], "Chooses axis bounds/ticks and converts values to canvas Y coordinates."),
            (["drawStandardBands(...) ", "trendChartGeometry(canvas)"], "Draws background categories and returns the chart's plotting rectangle."),
            (["escapeHtml(value)", "formatMetricTooltipValue(value, metric)", "comparisonValuesForDay(day)"], "Safely formats tooltip labels and collects same-day comparison values."),
            ("renderTrendChart()", "Draws axes, labels, standards bands, series lines, points, and legend data."),
            (["hideTrendTooltip()", "updateTrendTooltip(event)"], "Hides or repositions the interactive comparison tooltip."),
        ],
    )

    doc.add_heading("Block Group 6 — Report images and maps (app.js: 2455–3366)", level=1)
    add_block(
        doc,
        "Block 6.1 — Hero image canvas",
        "Draws the selected bundled/uploaded photo into each report hero canvas, with a generated fallback when no photo exists.",
        None,
        [
            ("renderScene(canvas)", "Crops and draws the selected photo into the report canvas."),
            ("drawSensorPictureFallback(ctx, width, height, fallback)", "Draws a text/shape placeholder for a sensor with no mapped image."),
        ],
    )
    add_block(
        doc,
        "Block 6.2 — Sensor selection and map values",
        "Chooses which catalog sensors appear, which are highlighted, and which reading/color each marker receives.",
        None,
        [
            (["sensorKind(sensor)", "mapColorForSensor(sensor)", "highlightedFillColor()"], "Resolves marker type and base/highlight colors."),
            (["mappedSensorsForMetric()", "selectedReportSensorIds()", "selectedMappedSensors()"], "Filters sensors by metric and active report selection."),
            (["sensorLocationValue(sensor)", "sensorRowsForSensor(sensor)", "snapshotSensorValue(sensor, metric)"], "Connects catalog sensors to loaded rows and computes the displayed snapshot value."),
            ("standardColorForValue(value, metric)", "Maps a snapshot reading to its standards-band color."),
            ("updateSensorMapDetails(sensors, selected)", "Updates the map title, legend, sensor list, and accessible description."),
        ],
    )
    add_block(
        doc,
        "Block 6.3 — Map geometry and drawing helpers",
        "Implements tile math, viewport fitting, collision-aware labels, markers, scale bars, and north arrows without a mapping framework.",
        """longitude/latitude -> Web Mercator world coordinates
  -> tile coordinates -> canvas coordinates -> marker/label""",
        [
            (["wait(ms)", "clamp(value, min, max)"], "General timing and bounds helpers."),
            (["lonToWorldX(...) ", "latToWorldY(...) ", "worldYToLat(...) ", "worldXToLon(...) ", "worldToTile(value)"], "Converts geographic coordinates, world pixels, and tile indexes."),
            ("staticMapViewport(sensors, width, height, focusSensors)", "Chooses zoom and center so relevant sensors fit the print map."),
            (["drawSensorPrintMapFallback(...) ", "loadTileImage(src)", "loadTileImageCached(src)"], "Loads raster tiles or draws a fallback background."),
            (["drawHexagonPath(...) ", "drawSensorMarker(...) ", "drawSnapshotSensorMarker(...)"], "Draws report markers."),
            (["rectanglesOverlap(...) ", "labelCandidateFits(...) ", "labelOverlapScore(...) ", "drawHighlightedLabel(...)"], "Places highlighted labels while minimizing collisions."),
            (["niceScaleDistance(...) ", "formatMetricScaleDistance(...) ", "formatMilesScaleDistance(...) ", "formatImperialScaleDistance(...)"], "Chooses and formats readable map scale lengths."),
            (["drawMapScaleBar(...) ", "drawNorthArrow(...) ", "mapTileUrl(...)"], "Draws map furniture and builds CARTO tile URLs."),
            (["snapshotPeriodLabel()", "drawSnapshotPeriodLabel(ctx)"], "Builds and draws the active day/month label."),
        ],
    )
    add_block(
        doc,
        "Block 6.4 — Map renderers and interactions",
        "Produces a static print map, optional interactive appendix map, and zoomable/pannable Template 3 snapshot map.",
        None,
        [
            ("renderStaticSensorPrintMap(sensors, selectedIds, focusSensors)", "Composes map tiles, markers, labels, scale bar, north arrow, and fallback behavior into the print canvas."),
            ("renderSensorMap()", "Shows/hides and updates the optional appendix map page."),
            (["canvasInternalPoint(event, canvas)", "ensureSnapshotMapView(...) ", "resizeCanvasForDisplay(canvas)"], "Converts pointer coordinates and maintains a correctly sized snapshot viewport."),
            ("renderSnapshotMap()", "Draws the interactive Template 3 sensor map and current reading colors."),
            ("setupSnapshotMapInteractions()", "Registers wheel/pinch zoom, drag pan, and reset controls."),
            (["loadPictureAsset(key)", "preloadPictureAssets()"], "Loads bundled report images before they are needed."),
        ],
    )

    doc.add_heading("Block Group 7 — Defaults and live sensor API (app.js: 3367–3814)", level=1)
    add_block(
        doc,
        "Block 7.1 — Dates, defaults, and sample startup",
        "Initializes report dates and sample rows so the interface is usable before a network request completes.",
        None,
        [
            (["plural(count, noun)", "formatReportDate(value)"], "Formats human-readable copy."),
            (["localDateParts(date)", "currentDateInputValue(date)", "currentMonthInputValue(date)", "currentReportDateLabel()"], "Produces timezone-safe input/report dates."),
            ("setCurrentPeriodDefaults()", "Initializes month, day, and report-date controls."),
            ("loadSampleData()", "Loads built-in demonstration rows and resets CSV/API state."),
            ("updateDataStatusForSelection()", "Explains whether the active selection has loaded readings."),
        ],
    )
    add_block(
        doc,
        "Block 7.2 — API request construction and normalization",
        "Maps the UI metric/location to the correct AQ or NU endpoint and converts the response back into the common row format.",
        """buildApiReadingsUrl(...)
  -> fetchWithTimeout(...)
  -> normalizeApiRows(...)
  -> mergeLoadedRows(...);""",
        [
            (["sensorApiConfigForReportMetric(metric)", "selectedLocationCanLoadSensorData(selection, metric)"], "Maps report metrics to API namespace/metric and validates the selection."),
            (["apiClusterId(selection)", "apiLocationId(selection)"], "Resolves numeric cluster or sensor API identifiers."),
            (["apiReadingDate(timestamp)", "apiReadingValue(reading, apiMetric)"], "Extracts normalized dates and values from flexible response shapes."),
            ("buildApiReadingsUrl(options)", "Builds individual or cluster daily-readings URLs and constrains aggregation to supported values."),
            (["nextIsoDate(dateString)", "sensorDataTimeoutMessage(timeoutMs)"], "Supports Template 3 one-day ranges and readable timeout errors."),
            ("fetchWithTimeout(resource, options)", "Wraps fetch with a per-request AbortController and timeout."),
            ("settleWithConcurrency(items, concurrency, task)", "Runs at most four sensor requests at once while retaining success/failure per item."),
            (["apiLoadedDateSpan(rows)", "blankMetricRow()"], "Summarizes returned dates and creates rows with null metrics."),
            ("normalizeApiRows(payload, context)", "Converts API readings into common rows tied to the selected location."),
            ("mergeLoadedRows(nextRows, context)", "Replaces only the requested metric/location/date slice and preserves other loaded data."),
        ],
    )
    add_block(
        doc,
        "Block 7.3 — Selection loading and orchestration",
        "Resolves one or many requested locations, executes the batch, accepts partial success, updates status, and triggers a render.",
        """loadApiData()
  -> selectedLoadLocations()
  -> settleWithConcurrency(fetchRowsForSelection)
  -> mergeLoadedRows()
  -> render();""",
        [
            ("fetchRowsForSelection(selection, options)", "Requests one sensor/cluster, handles the one-day end-date rule, filters returned dates, and normalizes rows."),
            (["loadLocationsForValues(values, metric)", "selectedLoadLocations(metric)"], "Resolves explicit builder members or the template's current locations."),
            (["syncCustomClusterLoadButtons()", "sensorTypeLabelForMetric(metric)"], "Keeps builder buttons and error text accurate."),
            ("loadApiData(options)", "Owns cancellation, batch timeout, concurrency, partial failures, merging, status messages, and final rendering."),
        ],
    )

    doc.add_heading("Block Group 8 — Report text, notes, templates, events, and photos (app.js: 3815–4838)", level=1)
    add_block(
        doc,
        "Block 8.1 — Comparison summaries and generated observations",
        "Turns calculated series and standards bands into the short explanatory text shown in summary cards and notes.",
        None,
        [
            ("updateComparisonSummary()", "Finds the highest peak, highest average, and coverage count for Template 2."),
            (["noteCanUseGeneratedText()", "formatMonthDay(value)", "formatTextList(items)"], "Determines when auto-copy is allowed and formats dates/lists."),
            (["rowsForGeneratedNote(metric)", "observationGroupsForGeneratedNote(rows, metric)"], "Collects note rows and groups them into standards categories."),
            (["generatedNoteObservationPhrase(group, totalCount)", "generatedNotePeriodLabel(info)"], "Builds readable category and period phrases."),
            (["buildAirQualityGeneratedNote(...) ", "buildHeatIndexGeneratedNote(...) ", "buildGeneratedNote(...)"], "Creates metric-specific observation text."),
            (["heatIndexGeneratedNoteHtml(note)", "generatedNoteHtml(note, metric)"], "Adds safe emphasis/markup to generated notes."),
            ("renderSnapshotObservationAreas(locationValues)", "Lists the locations represented on Template 3."),
        ],
    )
    add_block(
        doc,
        "Block 8.2 — Central text update and render loop",
        "Synchronizes report headings, metadata, labels, summary cards, standards, note contents, and all graphics after a state change.",
        """function render() {
  updateText(); renderCalendar(); renderScene(...);
  renderTrendChart(); renderSensorMap();
  state.snapshotMapPromise = renderSnapshotMap();
}""",
        [
            ("updateText()", "Writes selected dates, locations, author/catchphrase, summaries, generated notes, and standards into all three templates."),
            ("render()", "Single refresh entry point for text, calendar, images, trend chart, appendix map, and snapshot map."),
        ],
    )
    add_block(
        doc,
        "Block 8.3 — Rich note editor",
        "Keeps three contenteditable note areas synchronized, enforces a 100-word limit, and applies basic formatting commands.",
        None,
        [
            (["noteWords(text)", "truncateNoteText(text)", "enforceNoteWordLimit(editor)"], "Counts and limits note text."),
            ("moveCursorToEnd(node)", "Restores the caret after truncation."),
            (["syncNoteFromTextInput()", "syncNoteFromEditor(editor)"], "Copies note content between hidden input, state, and visible editors."),
            (["applyNoteCommand(editor, command, value)", "applyNoteFontSize(editor, size)"], "Applies bold/italic/underline/alignment/font formatting."),
            ("setupNoteEditor(config)", "Registers input, toolbar, font, size, and clear-format handlers."),
        ],
    )
    add_block(
        doc,
        "Block 8.4 — Template and control visibility",
        "Shows the correct controls/report panel and switches Template 3 between day and month modes.",
        None,
        [
            (["syncLocationModeControls()", "syncComparisonModeControls()"], "Shows the correct sensor/cluster/custom-builder controls."),
            (["configureAggregationOptions(template)", "syncSnapshotAggregationControls()"], "Provides month-only options for Templates 1/2 and day/month for Template 3."),
            ("setTemplate(template)", "Activates the chosen report panel, updates control visibility, and resets/renders the snapshot map."),
        ],
    )
    add_block(
        doc,
        "Block 8.5 — Event wiring and startup sequence",
        "The remaining top-level statements connect controls to the functions above. Most handlers update a DOM value or state field, call render(), and sometimes call loadApiData().",
        """setCurrentPeriodDefaults();
state.sensorCatalog = catalogKeyedByLocation(builtInSensorCatalog);
loadSampleData(); preloadPictureAssets(); setTemplate(state.template);
loadClusterCatalog().then(...); loadSensorCatalog().then(...);""",
        [
            ("Template/control listeners", "Handle template tabs, dates, metrics, locations, custom clusters, comparisons, searches, image upload, data loading, resizing, and PDF export."),
            ("PDF export handler", "Renders/waits for print maps, pauses briefly for paint completion, then calls window.print()."),
        ],
    )
    add_block(
        doc,
        "Block 8.6 — Photo registration and lookup",
        "Maps sensor selections and searchable labels to bundled photos or a deliberate no-photo fallback.",
        None,
        [
            ("registerBoxPhotoOptions()", "Adds deployment-box photos to the picture selector."),
            ("pictureKeyForSensorSelection(value)", "Resolves a selected sensor through API ID, built-in metadata, name aliases, and co-located box photos."),
            ("syncPictureToSensorSelection(value)", "Automatically selects the matching photo or configures the fallback canvas."),
            (["hidePictureSearchResults()", "selectPictureSearchResult(value, label)", "renderPictureSearchResults(matches)", "searchPictures(options)"], "Implements picture search and selection."),
        ],
    )

    doc.add_heading("Block Group 9 — HTML, CSS, local server, and diagnostics", level=1)
    add_block(
        doc,
        "Block 9.1 — index.html",
        "Defines controls and report output containers; JavaScript depends heavily on the element IDs.",
        """index.html
1–229    application header and input controls
230–336  Template 1: calendar report
337–428  Template 2: comparison trend report
429–512  Template 3: snapshot map report
513–541  optional sensor-map appendix""",
        [
            ("data-template-control", "Declares which controls belong to each template."),
            ("data-location-mode / data-comparison-mode", "Lets JavaScript show only the relevant sensor, cluster, or custom-builder controls."),
            ("Element IDs", "Form the contract consumed by the els object and all render functions."),
        ],
    )
    add_block(
        doc,
        "Block 9.2 — styles.css",
        "Separates screen-builder layout from printable report layout and styles each report component.",
        """styles.css
1–452     application shell and controls
453–907   report pages, headers, maps, shared panels
908–1285  screen-only report sizing, calendar, notes, charts
1286–1551 Template 3 and responsive rules
1552–2033 print/PDF rules""",
        [
            ("Screen rules", "Lay out the left control panel and scaled report preview."),
            ("Component rules", "Style calendar cells, legends, charts, maps, note editors, photos, and summaries."),
            ("Print rules", "Hide the builder UI and size the selected report plus optional appendix for browser PDF output."),
        ],
    )
    add_block(
        doc,
        "Block 9.3 — local-server.mjs",
        "Serves local static files and provides protected proxy routes. The current main report reads live results directly from the Common Senses API; the QuantAQ proxy is retained but not used by app.js.",
        None,
        [
            ("sendJson(response, status, payload)", "Sends non-cached JSON errors/responses."),
            ("readLocalConfig()", "Reads config.local.json without exposing it as a static file."),
            ("handleQuantaqReadings(requestUrl, response)", "Validates the local API key and proxies a QuantAQ resampled-data request."),
            ("handleClusterList(response)", "Proxies the Common Senses cluster catalog for local fallback use."),
            ("staticFilePath(pathname)", "Normalizes paths and rejects traversal outside the project root."),
            ("serveStatic(requestUrl, response)", "Blocks dotfiles/local config and serves recognized files with MIME types."),
            ("createServer(...) routing", "Allows only GET, dispatches proxy routes, otherwise serves static content."),
        ],
    )
    add_block(
        doc,
        "Block 9.4 — internal/load-diagnostics.js",
        "Runs repeated uncached API requests locally, records timings in browser storage, draws a response-time chart, and exports/prints results.",
        None,
        [
            (["localDateValue(date)", "setDateDefaults()", "currentSettings()", "validateSettings(settings)"], "Initializes and validates diagnostic inputs."),
            (["readHistory()", "storeHistory()"], "Loads/saves the last 40 attempts in localStorage."),
            (["syncMetricOptions()", "durationLabel(durationMs)", "requestEndpoint(settings)", "buildUrl(settings, attemptNumber)"], "Aligns pipeline/metric choices and builds uncached diagnostic URLs."),
            ("runAttempt(settings, sequence, parentSignal)", "Times one request and records success, timeout, HTTP status, row count, and error."),
            (["renderChart()", "renderTable()", "updatePrintDetails()", "render()"], "Draws and updates the diagnostic results and print view."),
            (["runDiagnostic()", "csvCell(value)", "downloadCsv()"], "Runs attempts sequentially, supports Stop, and exports recorded results as escaped CSV."),
        ],
    )

    doc.core_properties.title = "CSENSES Report Builder: Code Blocks and Functions"
    doc.core_properties.subject = "Code-first explanation organized by block groups, blocks, and functions"
    doc.core_properties.author = "OpenAI Codex"
    doc.core_properties.keywords = "CSENSES, JavaScript, functions, code blocks, report builder"
    doc.core_properties.comments = "Revised from source at commit e15cf32 on 2026-09-23."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
